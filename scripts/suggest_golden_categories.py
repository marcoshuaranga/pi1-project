"""Genera sugerencias de categoría top-9 a partir de las etiquetas en texto libre de los expertos.

Uso:
    uv run python scripts/suggest_golden_categories.py
Escribe data/processed/golden_annotation/sugerencias_categoria.xlsx para que los expertos confirmen.
No modifica experto1.xlsx ni experto2.xlsx.
"""

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from pi_core.config import get_settings
from pi_core.constants import GOLDEN_CATEGORIES
from pi_core.evaluation.golden_set import ANNOTATION_DIR, EXPERT1_FILE, EXPERT2_FILE
from pi_core.evaluation.golden_suggest import suggest_category

OUT_FILE = "sugerencias_categoria.xlsx"


def _read_raw(path: Path) -> dict[str, str]:
    ws = load_workbook(path, data_only=True)["anotacion"]
    return {str(r[0]): str(r[2]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}


def main() -> None:
    processed = Path(get_settings().data_processed_path)
    ann = processed / ANNOTATION_DIR
    e1, e2 = _read_raw(ann / EXPERT1_FILE), _read_raw(ann / EXPERT2_FILE)
    pool = json.loads((processed / "golden_pool.json").read_text(encoding="utf-8"))
    evals = {t["ticket_id"]: t for t in json.loads((processed / "tickets_eval.json").read_text(encoding="utf-8"))}

    wb = Workbook()
    ws = wb.active
    ws.title = "sugerencias"
    ws.append(
        ["ticket_id", "titulo_anon", "etiqueta_experto1", "etiqueta_experto2", "categoria_sugerida",
         "confianza", "motivo", "categoria_confirmada (elegir de la lista)"]
    )
    stats: Counter = Counter()
    for e in pool:
        tid = e["ticket_id"]
        sol = evals.get(tid, {}).get("solucion_anon") or ""
        cat, conf, motivo = suggest_category(e1[tid], e["titulo_anon"], sol)
        stats[(conf, cat is not None)] += 1
        ws.append([tid, e["titulo_anon"], e1[tid], e2.get(tid), cat, conf, motivo, None])

    cats = wb.create_sheet("categorias")
    for c in GOLDEN_CATEGORIES:
        cats.append([c])
    cats.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1=f"=categorias!$A$1:$A${len(GOLDEN_CATEGORIES)}", allow_blank=True)
    dv.add(f"H2:H{len(pool) + 1}")
    ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGH", (14, 70, 18, 18, 62, 11, 55, 62), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    out = ann / OUT_FILE
    wb.save(out)
    print(f"Sugerencias: {out}")
    for (conf, has_cat), n in sorted(stats.items()):
        print(f"  confianza={conf:5} {'con sugerencia' if has_cat else 'SIN sugerencia'}: {n}")


if __name__ == "__main__":
    main()
