"""C9 — Golden set: muestreo estratificado + doble anotación ciega (opción B).

Diseño (ver DEMO_QA.md "¿Qué es el golden set?"): 200 tickets estratificados por
`categoria_top9`, de los cuales un subconjunto (por defecto 30) se etiqueta por
partida doble para medir el acuerdo inter-anotador (Cohen's Kappa); el resto se
etiqueta una sola vez. Los dos expertos anotan a ciegas — nunca ven la
`categoria_top9` original ni la respuesta del otro experto — y solo reciben
`titulo_anon`, el mismo insumo que usa `ClassifierAgent.classify` en
`EvaluationFramework`, para que la comparación F1 sea justa.

Flujo:
    1. `build_pool()`      → golden_pool.json (interno) + planillas .xlsx en blanco
    2. (fuera de este módulo) 2 expertos llenan las planillas, a ciegas
    3. `score_and_merge()` → golden_set_sample.json + golden_set_kappa.json

Ver scripts/build_golden_set.py para el CLI.
"""

from __future__ import annotations

import json
import random
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import cohen_kappa_score

from pi_core.constants import TOP9_CATEGORIES

POOL_FILE = "golden_pool.json"
GOLDEN_FILE = "golden_set_sample.json"
KAPPA_FILE = "golden_set_kappa.json"
ANNOTATION_DIR = "golden_annotation"
EXPERT1_FILE = "experto1.xlsx"
EXPERT2_FILE = "experto2.xlsx"

_CATEGORIA_HEADER = "categoria (elegir de la lista desplegable)"
_SHEET_NAME = "anotacion"


def _stratified_quota(counts: dict[str, int], n: int) -> dict[str, int]:
    """Asignación proporcional (método del resto mayor) de `n` entre categorías."""
    total = sum(counts.values())
    if total <= n:
        return dict(counts)

    exact = {cat: n * count / total for cat, count in counts.items()}
    quota = {cat: min(counts[cat], int(v)) for cat, v in exact.items()}
    remaining = n - sum(quota.values())

    by_remainder = sorted(counts, key=lambda c: exact[c] - quota[c], reverse=True)
    for cat in by_remainder:
        if remaining <= 0:
            break
        if quota[cat] < counts[cat]:
            quota[cat] += 1
            remaining -= 1
    return quota


def sample_golden_pool(
    processed: Path,
    seed: int,
    n: int = 200,
    n_kappa: int = 30,
    eval_filename: str = "tickets_eval.json",
) -> list[dict[str, Any]]:
    """Muestra estratificada por `categoria_top9`, con `n_kappa` marcados para doble anotación."""
    eval_path = processed / eval_filename
    if not eval_path.exists():
        raise FileNotFoundError(f"No se encontró {eval_path}; corre el pipeline (extract→ingest) primero")

    tickets = json.loads(eval_path.read_text(encoding="utf-8"))
    candidates = [
        t for t in tickets if t.get("categoria_top9") in TOP9_CATEGORIES and t.get("titulo_anon")
    ]
    if not candidates:
        raise ValueError(f"{eval_path} no tiene tickets con categoria_top9 válida")

    by_cat: dict[str, list[dict[str, Any]]] = {}
    for t in candidates:
        by_cat.setdefault(t["categoria_top9"], []).append(t)

    quota = _stratified_quota({cat: len(items) for cat, items in by_cat.items()}, n)

    rng = random.Random(seed)
    pool: list[dict[str, Any]] = []
    for cat, items in by_cat.items():
        pool.extend(rng.sample(items, quota.get(cat, 0)))
    rng.shuffle(pool)

    kappa_ids = {t["ticket_id"] for t in pool[:n_kappa]}
    return [
        {
            "golden_id": str(uuid.uuid4()),
            "ticket_id": t["ticket_id"],
            "titulo_anon": t["titulo_anon"],
            # Solo para auditoría interna — nunca se exporta a las planillas de los expertos.
            "categoria_top9": t["categoria_top9"],
            "doble_anotacion": t["ticket_id"] in kappa_ids,
        }
        for t in pool
    ]


def _new_annotation_workbook(rows: list[dict[str, Any]], shuffle_seed: int) -> Workbook:
    order = list(rows)
    random.Random(shuffle_seed).shuffle(order)

    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET_NAME
    ws.append(["ticket_id", "titulo_anon", _CATEGORIA_HEADER])
    for row in order:
        ws.append([row["ticket_id"], row["titulo_anon"], ""])

    cats_ws = wb.create_sheet("categorias")
    for cat in TOP9_CATEGORIES:
        cats_ws.append([cat])
    cats_ws.sheet_state = "hidden"

    last_row = len(order) + 1
    dv = DataValidation(
        type="list",
        formula1=f"=categorias!$A$1:$A${len(TOP9_CATEGORIES)}",
        allow_blank=True,
    )
    dv.error = "Elige una categoría de la lista desplegable"
    dv.errorTitle = "Categoría inválida"
    dv.add(f"C2:C{last_row}")
    ws.add_data_validation(dv)

    for col, width in (("A", 14), ("B", 90), ("C", 55)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return wb


def build_pool(processed: Path, seed: int, n: int = 200, n_kappa: int = 30) -> list[dict[str, Any]]:
    """Genera golden_pool.json y las 2 planillas .xlsx en blanco para los expertos."""
    entries = sample_golden_pool(processed, seed, n=n, n_kappa=n_kappa)

    processed.mkdir(parents=True, exist_ok=True)
    (processed / POOL_FILE).write_text(json.dumps(entries, indent=2, ensure_ascii=False), encoding="utf-8")

    ann_dir = processed / ANNOTATION_DIR
    ann_dir.mkdir(parents=True, exist_ok=True)

    rows = [{"ticket_id": e["ticket_id"], "titulo_anon": e["titulo_anon"]} for e in entries]
    kappa_rows = [r for r, e in zip(rows, entries, strict=True) if e["doble_anotacion"]]

    _new_annotation_workbook(rows, shuffle_seed=seed).save(ann_dir / EXPERT1_FILE)
    _new_annotation_workbook(kappa_rows, shuffle_seed=seed + 1).save(ann_dir / EXPERT2_FILE)

    return entries


def _read_annotations(path: Path) -> dict[str, str | None]:
    if not path.exists():
        raise FileNotFoundError(f"No se encontró la planilla llena: {path}")

    wb = load_workbook(path, data_only=True)
    ws = wb[_SHEET_NAME] if _SHEET_NAME in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    try:
        id_col = header.index("ticket_id")
        cat_col = header.index(_CATEGORIA_HEADER)
    except ValueError as exc:
        raise ValueError(f"{path}: encabezados inesperados {header}") from exc

    out: dict[str, str | None] = {}
    for row in rows:
        ticket_id = row[id_col]
        if ticket_id is None:
            continue
        ticket_id = str(ticket_id)
        categoria = row[cat_col]
        if categoria:
            categoria = str(categoria).strip()
            if categoria not in TOP9_CATEGORIES:
                raise ValueError(f"{path}: '{categoria}' (ticket {ticket_id}) no es una categoría top-9 válida")
        out[ticket_id] = categoria or None
    return out


def score_and_merge(processed: Path) -> dict[str, Any]:
    """Lee las 2 planillas llenas, calcula Kappa y escribe golden_set_sample.json.

    Desacuerdos en el subconjunto de doble anotación quedan con
    `categoria_gold: null` y `necesita_arbitraje: true` — se resuelven a mano
    (tercer árbitro o discusión) antes de usar el golden set para evaluar.
    """
    pool_path = processed / POOL_FILE
    if not pool_path.exists():
        raise FileNotFoundError(f"No se encontró {pool_path}; corre build_pool()/`sample` primero")
    pool = json.loads(pool_path.read_text(encoding="utf-8"))

    experto1 = _read_annotations(processed / ANNOTATION_DIR / EXPERT1_FILE)
    experto2 = _read_annotations(processed / ANNOTATION_DIR / EXPERT2_FILE)

    missing1 = [e["ticket_id"] for e in pool if not experto1.get(e["ticket_id"])]
    if missing1:
        raise ValueError(f"Experto 1: faltan {len(missing1)} categorías sin llenar: {missing1[:5]}")

    kappa_ids = [e["ticket_id"] for e in pool if e["doble_anotacion"]]
    missing2 = [tid for tid in kappa_ids if not experto2.get(tid)]
    if missing2:
        raise ValueError(f"Experto 2: faltan {len(missing2)} categorías sin llenar: {missing2[:5]}")

    y1 = [experto1[tid] for tid in kappa_ids]
    y2 = [experto2[tid] for tid in kappa_ids]
    kappa = cohen_kappa_score(y1, y2) if kappa_ids else None

    entries = []
    n_arbitraje = 0
    for e in pool:
        tid = e["ticket_id"]
        cat1 = experto1[tid]
        cat2 = experto2.get(tid)
        acuerdo = None
        gold = cat1
        necesita_arbitraje = False
        if e["doble_anotacion"]:
            acuerdo = cat1 == cat2
            if not acuerdo:
                necesita_arbitraje = True
                gold = None
                n_arbitraje += 1
        entries.append(
            {
                "golden_id": e["golden_id"],
                "ticket_id": tid,
                "categoria_experto1": cat1,
                "categoria_experto2": cat2,
                "categoria_gold": gold,
                "doble_anotacion": e["doble_anotacion"],
                "acuerdo": acuerdo,
                "necesita_arbitraje": necesita_arbitraje,
                "split": "eval",
            }
        )

    golden_path = processed / GOLDEN_FILE
    golden_path.write_text(json.dumps(entries, indent=2, ensure_ascii=False), encoding="utf-8")

    kappa_summary = {
        "kappa": round(kappa, 4) if kappa is not None else None,
        "n_doble_anotado": len(kappa_ids),
        "n_total": len(entries),
        "n_arbitraje_pendiente": n_arbitraje,
        "fecha_calculo": datetime.now(UTC).isoformat(),
    }
    (processed / KAPPA_FILE).write_text(
        json.dumps(kappa_summary, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    return {"golden_path": str(golden_path), "kappa_path": str(processed / KAPPA_FILE), **kappa_summary}
