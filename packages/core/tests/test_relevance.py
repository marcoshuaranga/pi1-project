"""Juicio experto de relevancia: pool a ciegas, planillas y Recall@5 / Precisión@5 / Kappa."""

import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from pi_core.constants import FUERA_TOP9, TOP9_CATEGORIES
from pi_core.evaluation.golden_set import ANNOTATION_DIR
from pi_core.evaluation.relevance import (
    EXPERT1_FILE,
    EXPERT2_FILE,
    build_relevance_pool,
    score_relevance,
)

CAT = TOP9_CATEGORIES[0]


class FakeRag:
    top_k = 5

    def retrieve(self, texto, categoria=None):
        assert self.top_k == 6  # un extra para poder descartar el propio ticket
        propio = texto.replace("titulo ", "T")
        ids = [propio, "r1", "r2", "r3", "r4", "r5"]
        return [
            SimpleNamespace(articulo_o_ticket_id=i, tipo="ticket", titulo=f"tit {i}", solucion=f"sol {i}", score=0.9)
            for i in ids
        ]


def _setup(tmp_path, n_tickets=4):
    golden, evals = [], []
    for i in range(n_tickets):
        tid = f"T{i}"
        golden.append(
            {"ticket_id": tid, "categoria_gold": CAT, "doble_anotacion": i < 3}
        )
        evals.append({"ticket_id": tid, "titulo_anon": f"titulo {i}", "solucion_anon": f"real {i}"})
    golden.append({"ticket_id": "TX", "categoria_gold": FUERA_TOP9, "doble_anotacion": True})
    evals.append({"ticket_id": "TX", "titulo_anon": "titulo X", "solucion_anon": ""})
    (tmp_path / "golden_set_sample.json").write_text(json.dumps(golden), encoding="utf-8")
    (tmp_path / "tickets_eval.json").write_text(json.dumps(evals), encoding="utf-8")


def _fill(path, decide):
    wb = load_workbook(path)
    ws = wb["relevancia"]
    for row in ws.iter_rows(min_row=2):
        row[5].value = decide(row[0].value)
    wb.save(path)


def test_pool_excluye_fuera_top9_propio_ticket_y_no_filtra_ranking(tmp_path):
    _setup(tmp_path)
    items = build_relevance_pool(tmp_path, FakeRag(), n=3, k=5)

    assert [i["ticket_id"] for i in items] == ["T0", "T1", "T2"]  # doble anotación primero, sin "fuera"
    assert all(len(i["candidatas"]) == 5 for i in items)
    assert all(c["id"] != i["ticket_id"] for i in items for c in i["candidatas"])

    ws = load_workbook(tmp_path / ANNOTATION_DIR / EXPERT1_FILE)["relevancia"]
    header = [c.value for c in ws[1]]
    assert header[0] == "item_id" and "score" not in " ".join(map(str, header)).lower()
    assert "categoria" not in " ".join(map(str, header)).lower()
    assert ws.max_row == 1 + 3 * 5

    orden1 = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
    orden2 = [r[0] for r in load_workbook(tmp_path / ANNOTATION_DIR / EXPERT2_FILE)["relevancia"].iter_rows(min_row=2, values_only=True)]
    assert sorted(orden1) == sorted(orden2) and orden1 != orden2  # mismo contenido, distinto orden


def test_score_relevance_calcula_recall_precision_y_acuerdo(tmp_path):
    _setup(tmp_path)
    build_relevance_pool(tmp_path, FakeRag(), n=3, k=5)
    ann = tmp_path / ANNOTATION_DIR

    # Experto 1: solo el ticket T0 tiene algo útil (su candidata #1). Experto 2: igual, más T1#2.
    _fill(ann / EXPERT1_FILE, lambda i: "Sí" if i == "T0#1" else "No")
    _fill(ann / EXPERT2_FILE, lambda i: "Sí" if i in ("T0#1", "T1#2") else "No")

    r = score_relevance(tmp_path)

    assert r["n_tickets"] == 3 and r["n_juicios"] == 15
    assert r["recall_at_5_experto"] == pytest.approx(1 / 3, abs=1e-4)  # conservador: ambos Sí -> solo T0
    assert r["recall_at_5_experto_laxo"] == pytest.approx(2 / 3, abs=1e-4)
    assert r["precision_at_5_experto"] == pytest.approx(1 / 15, abs=1e-4)
    assert r["por_experto"]["experto1"]["recall_at_5"] == pytest.approx(1 / 3, abs=1e-4)
    assert r["por_experto"]["experto2"]["recall_at_5"] == pytest.approx(2 / 3, abs=1e-4)
    assert r["acuerdo_relevancia"] == pytest.approx(14 / 15, abs=1e-4)
    assert (tmp_path / "relevance_results.json").exists()


def test_score_relevance_rechaza_respuestas_incompletas_o_invalidas(tmp_path):
    _setup(tmp_path)
    build_relevance_pool(tmp_path, FakeRag(), n=2, k=5)
    ann = tmp_path / ANNOTATION_DIR

    with pytest.raises(ValueError, match="respuesta inválida"):
        score_relevance(tmp_path)  # planillas en blanco

    _fill(ann / EXPERT1_FILE, lambda i: "tal vez")
    with pytest.raises(ValueError, match="respuesta inválida"):
        score_relevance(tmp_path)


def test_golden_evaluation_incluye_recall_experto_si_existe(tmp_path):
    from pi_core.evaluation.metrics import EvaluationFramework

    ev = EvaluationFramework.__new__(EvaluationFramework)
    ev.processed = tmp_path
    assert ev._load_relevance_experto() == {}

    (tmp_path / "relevance_results.json").write_text(
        json.dumps({"recall_at_5_experto": 0.7, "precision_at_5_experto": 0.4, "n_tickets": 30, "kappa_relevancia": 0.6}),
        encoding="utf-8",
    )
    assert ev._load_relevance_experto() == {
        "recall_at_5_experto": 0.7,
        "precision_at_5_experto": 0.4,
        "recall_at_5_experto_n": 30,
        "kappa_relevancia": 0.6,
    }
